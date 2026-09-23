"""Independent workbook fixtures; real partner reports are never modified."""

from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from hackalem.importers.systeme import parse_workbook


XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
TX_HEADERS = [
    "Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество",
]
MONTHS = ["Янв", "Фев", "Мар", "Апр", "Май", "Июн", "Июл", "Авг", "Сен", "Окт", "Ноя", "Дек"]


def _transactions_file(path: Path, quantities: list) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Детализация"
    sheet.append(TX_HEADERS)
    for quantity in quantities:
        sheet.append([
            "22.09.2026 16:10:12", "000012345", "Расходная накладная",
            "0007_", "Тестовый товар", "шт", "Алматы", quantity,
        ])
    workbook.save(path)
    workbook.close()
    return path


def _rewrite_fixture_sheet(path: Path, transform, sheet_number: int = 1) -> None:
    """Change synthetic XML fields that openpyxl cannot save directly."""
    member = f"xl/worksheets/sheet{sheet_number}.xml"
    with ZipFile(path) as archive:
        contents = [(info, archive.read(info.filename)) for info in archive.infolist()]
    replacement = path.with_suffix(".rewritten.xlsx")
    with ZipFile(replacement, "w") as archive:
        for info, content in contents:
            if info.filename == member:
                xml = ET.fromstring(content)
                transform(xml)
                content = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
            archive.writestr(info, content)
    replacement.replace(path)


def _set_formula_cache(path: Path, coordinate: str, value: str) -> None:
    def transform(xml):
        cell = xml.find(f".//{{{XML_NS}}}c[@r='{coordinate}']")
        assert cell is not None
        cached = cell.find(f"{{{XML_NS}}}v")
        if cached is None:
            cached = ET.SubElement(cell, f"{{{XML_NS}}}v")
        cached.text = value

    _rewrite_fixture_sheet(path, transform)


def test_transaction_rows_preserve_duplicates_codes_and_quantity_states(tmp_path):
    path = _transactions_file(tmp_path / "source.xlsx", [5, 5, -4, None, 0, "#N/A"])
    source_bytes = path.read_bytes()

    parsed = parse_workbook(path)

    assert parsed.source_kind == "transactions"
    assert len(parsed.transactions) == 6
    assert [item["row"] for item in parsed.transactions] == [2, 3, 4, 5, 6, 7]
    assert [item["quantity"] for item in parsed.transactions] == [5, 5, -4, None, 0, None]
    assert [item["state"] for item in parsed.transactions] == [
        "value", "value", "value", "blank", "value", "error",
    ]
    assert {item["sku"] for item in parsed.transactions} == {"0007_"}
    assert {item["document_number"] for item in parsed.transactions} == {"000012345"}
    assert {item["sku"] for item in parsed.products} == {"0007_"}
    assert any(issue["code"] == "NEGATIVE_QUANTITY" for issue in parsed.issues)
    assert path.read_bytes() == source_bytes


def test_invalid_declared_xml_dimension_does_not_truncate_transactions(tmp_path):
    path = _transactions_file(tmp_path / "wrong-dimension.xlsx", [2, 3, 4])

    def transform(xml):
        dimension = xml.find(f"{{{XML_NS}}}dimension")
        assert dimension is not None
        dimension.set("ref", "H4:H4")

    _rewrite_fixture_sheet(path, transform)

    parsed = parse_workbook(path)

    assert [item["quantity"] for item in parsed.transactions] == [2, 3, 4]
    assert [item["cell"] for item in parsed.transactions] == ["H2", "H3", "H4"]
    sheet = next(item for item in parsed.sheets if item["sheet"] == "Детализация")
    assert sheet["declared_dimension"] == "H4:H4"
    assert sheet["max_row"] == 4
    assert sheet["max_column"] == 8


def test_headers_determine_source_kind_even_with_misleading_filename(tmp_path):
    path = _transactions_file(tmp_path / "MOQ SystemElectric.xlsx", [7])

    parsed = parse_workbook(path)

    assert parsed.source_kind == "transactions"
    assert len(parsed.transactions) == 1
    assert parsed.transactions[0]["quantity"] == 7


def test_familiar_filename_cannot_replace_required_headers(tmp_path):
    workbook = Workbook()
    workbook.active.append(["Произвольное поле", "Другое поле"])
    workbook.active.append(["0007_", 9])
    path = tmp_path / "Динамика продаж_Syseme Electric_2025-2026.xlsx"
    workbook.save(path)
    workbook.close()

    with pytest.raises(ValueError):
        parse_workbook(path)


@pytest.mark.parametrize("cached_value", [None, "5"])
def test_formula_keeps_expression_and_uses_only_saved_cache(tmp_path, cached_value):
    path = _transactions_file(tmp_path / "formula.xlsx", ["=2+3"])
    if cached_value is not None:
        _set_formula_cache(path, "H2", cached_value)

    parsed = parse_workbook(path)

    raw = next(item for item in parsed.raw_rows if item["row"] == 2)
    assert raw["cells"]["H"]["formula"] == "=2+3"
    transaction = parsed.transactions[0]
    if cached_value is None:
        assert raw["cells"]["H"]["value"] is None
        assert transaction["quantity"] is None
        assert transaction["state"] == "error"
        assert any(issue["code"] == "FORMULA_CACHE_MISSING" for issue in parsed.issues)
    else:
        assert raw["cells"]["H"]["value"] == 5
        assert transaction["quantity"] == 5
        assert transaction["state"] == "value"
        assert not any(issue["code"] == "FORMULA_CACHE_MISSING" for issue in parsed.issues)


def test_hidden_seasonality_sheet_preserves_cached_observation_and_blank_month(tmp_path):
    workbook = Workbook()
    main = workbook.active
    main.title = "Товары"
    main.append(["Номенклатура", "Номенклатура.Код", "Артикул", "Кратность"])
    main.append(["Тестовый товар", "0007_", "ART0007", 6])
    hidden = workbook.create_sheet("Сезонный расчёт")
    hidden.sheet_state = "hidden"
    hidden.append(["Год", *MONTHS])
    hidden.append([2025, "=2+3", *range(2, 12), None])
    path = tmp_path / "multiple-and-hidden-season.xlsx"
    workbook.save(path)
    workbook.close()

    def cache_hidden_formula(xml):
        cell = xml.find(f".//{{{XML_NS}}}c[@r='B2']")
        assert cell is not None
        value = cell.find(f"{{{XML_NS}}}v")
        if value is None:
            value = ET.SubElement(cell, f"{{{XML_NS}}}v")
        value.text = "5"

    _rewrite_fixture_sheet(path, cache_hidden_formula, sheet_number=2)

    parsed = parse_workbook(path)

    sheet = next(item for item in parsed.sheets if item["sheet"] == hidden.title)
    assert sheet["state"] == "hidden"
    observations = {
        item["period"]: item for item in parsed.seasonal_values
        if item["sheet"] == hidden.title
    }
    assert len(observations) == 12
    assert observations["2025-01-01"]["value"] == 5
    assert observations["2025-01-01"]["state"] == "value"
    assert observations["2025-12-01"]["value"] is None
    assert observations["2025-12-01"]["state"] == "blank"
    raw = next(item for item in parsed.raw_rows if item["sheet"] == hidden.title and item["row"] == 2)
    assert raw["cells"]["B"]["formula"] == "=2+3"
    assert raw["cells"]["B"]["value"] == 5
