"""Stage 11: stable corrections, immutable approvals and verified exports."""

import csv
import io
import json
import shutil
import sqlite3
from pathlib import Path

import pytest
from openpyxl import load_workbook

from hackalem.__main__ import main
from hackalem.services.cleaning import run_cleaning
from hackalem.services.forecasting import run_forecast
from hackalem.services.imports import list_snapshots
from hackalem.services.orders import (
    approve_order, build_order_export, create_order_project, export_order_file,
    list_order_versions, order_report, submit_order_for_review, update_order_item,
)
from hackalem.services.replenishment import replenishment_report, run_replenishment
from hackalem.services.quality import list_quality_runs
from hackalem.services.synthetic import create_synthetic_dataset
from hackalem.storage import SCHEMA_VERSION, initialize_database


def _forecast_configuration():
    return {
        "horizon_months": 12, "warehouse_scope": "source_report",
        "growth_application": {
            "mode": "replace_trend", "start": "2026-01-01", "end": "2026-12-01",
            "scope": "source_report", "status": "scenario",
            "reason": "Проверка заказов", "author": "pytest",
        },
        "short_history_fallback": None,
        "seasonal_aggregate_policy": {
            "use": False, "unit_status": "unknown",
            "reason": "Единица не подтверждена", "author": "pytest",
        },
    }


@pytest.fixture(scope="module")
def order_source(tmp_path_factory):
    root = tmp_path_factory.mktemp("orders")
    generated = create_synthetic_dataset(root / "datasets", seed=20260923)
    database = Path(generated["database_path"])
    snapshot = next(row for row in generated["snapshots"]
                    if row["supplier"] == "Systeme Electric")
    cleaning = run_cleaning(database, snapshot["snapshot_id"], "2026-01-01")
    forecast = run_forecast(
        database, snapshot["run_id"], cleaning["run_id"], "SYN-A-001",
        _forecast_configuration(), allow_scenario=True,
    )
    payload = {
        "quality_run_id": snapshot["run_id"], "cleaning_run_id": cleaning["run_id"],
        "as_of": "2026-01-01", "supplier": "Systeme Electric",
        "warehouse": "all_selected_warehouses", "items": [{
            "sku": "SYN-A-001", "category_code": "SYN-CAT-01",
            "forecast_run_id": forecast["run_id"], "project_commitments": [],
        }],
    }
    replenishment = run_replenishment(database, payload)
    return database, snapshot, replenishment


def _copy_database(order_source, tmp_path):
    source, snapshot, replenishment = order_source
    target = tmp_path / "hackalem.sqlite3"
    shutil.copyfile(source, target)
    return target, snapshot, replenishment


def _approved_order(database, run_id):
    draft = create_order_project(database, run_id, "Менеджер")
    item = draft["items"][0]
    corrected = update_order_item(
        database, draft["version_id"], item["sku"],
        item["selected_quantity"] + 12, "Менеджер", "Ручная проверка спроса",
    )
    review = submit_order_for_review(
        database, corrected["version_id"], "Менеджер", "Строки проверены",
    )
    return approve_order(
        database, review["version_id"], "Ответственный", "Локальное утверждение теста",
    )


def test_correction_survives_reload_and_approved_version_is_immutable(order_source, tmp_path):
    database, snapshot, replenishment = _copy_database(order_source, tmp_path)
    draft = create_order_project(database, replenishment["run_id"], "Менеджер")
    item = draft["items"][0]
    suggested = item["suggested_quantity"]
    changed = update_order_item(
        database, draft["version_id"], item["sku"], suggested + 12,
        "Менеджер", "Проверен клиентский проект",
    )
    reloaded = order_report(database, changed["version_id"])
    assert reloaded["items"][0]["line_id"] == item["sku"]
    assert reloaded["items"][0]["suggested_quantity"] == suggested
    assert reloaded["items"][0]["selected_quantity"] == suggested + 12
    assert reloaded["items"][0]["correction_reason"] == "Проверен клиентский проект"
    assert list_order_versions(database, snapshot["snapshot_id"])[0]["version_id"] == changed["version_id"]

    review = submit_order_for_review(database, changed["version_id"], "Менеджер", "Передано старшему")
    approved = approve_order(database, review["version_id"], "Старший менеджер", "Количество и единицы проверены")
    frozen = json.loads(json.dumps(approved["approval_snapshot"], ensure_ascii=False))
    assert approved["status_label"] == "Утверждён"
    assert frozen["responsible_is_authenticated_identity"] is False

    revision = update_order_item(
        database, approved["version_id"], item["sku"], suggested + 24,
        "Менеджер", "Новая заявка клиента",
    )
    assert revision["status_label"] == "Черновик"
    assert revision["parent_version_id"] == approved["version_id"]
    assert revision["version_number"] == approved["version_number"] + 1
    unchanged = order_report(database, approved["version_id"])
    assert unchanged["approval_snapshot"] == frozen
    assert unchanged["items"][0]["selected_quantity"] == suggested + 12


def test_real_scenario_cannot_be_approved(order_source, tmp_path):
    database, _, replenishment = _copy_database(order_source, tmp_path)
    with sqlite3.connect(database) as connection:
        connection.execute("INSERT OR REPLACE INTO app_metadata VALUES ('dataset_kind','real')")
        connection.execute("INSERT OR REPLACE INTO app_metadata VALUES ('dataset_id','real-test')")
    payload = replenishment_report(database, replenishment["run_id"])["input"]["payload"]
    real_scenario = run_replenishment(database, payload)
    assert real_scenario["input"]["dataset"]["kind"] == "real"
    draft = create_order_project(database, real_scenario["run_id"], "Менеджер")
    review = submit_order_for_review(database, draft["version_id"], "Менеджер", "Проверить сценарий")
    with pytest.raises(ValueError, match="сценарные входы"):
        approve_order(database, review["version_id"], "Ответственный", "Попытка утверждения")
    assert order_report(database, review["version_id"])["status"] == "review"


def test_csv_and_xlsx_exports_match_approved_version(order_source, tmp_path):
    database, _, replenishment = _copy_database(order_source, tmp_path)
    approved = _approved_order(database, replenishment["run_id"])
    selected = approved["items"][0]["selected_quantity"]
    unit = approved["items"][0]["purchase_unit"]

    csv_export = build_order_export(database, approved["version_id"], "csv")
    assert csv_export["verified"] is True
    assert csv_export["metadata"]["classification"] == "SYNTHETIC_SCENARIO"
    decoded = csv_export["content"].decode("utf-8-sig")
    csv_rows = list(csv.DictReader(io.StringIO(decoded), delimiter=";"))
    assert len(csv_rows) == 1
    assert float(csv_rows[0]["Количество"]) == selected
    assert csv_rows[0]["Единица закупки"] == unit
    assert csv_rows[0]["Код1С"] == "SYN-A-001"
    assert "Ручная проверка спроса" in decoded

    xlsx_export = build_order_export(database, approved["version_id"], "xlsx")
    assert xlsx_export["verified"] is True
    workbook = load_workbook(io.BytesIO(xlsx_export["content"]), read_only=True)
    values = list(workbook["Заказ"].iter_rows(values_only=True))
    row = dict(zip(values[0], values[1]))
    assert row["Количество"] == selected
    assert row["Единица закупки"] == unit
    assert row["Код1С"] == "SYN-A-001"
    assert row["Маркировка"] == "SYNTHETIC_SCENARIO"

    saved = export_order_file(database, approved["version_id"], tmp_path / "заказ.xlsx")
    assert saved["verified"] is True and saved["size"] > 0
    assert saved["rows"] == xlsx_export["rows"]
    saved_book = load_workbook(saved["path"], read_only=True)
    assert saved_book["Заказ"]["F2"].value == selected


def test_order_cli_reports_and_exports_selected_version(order_source, tmp_path, monkeypatch, capsys):
    database, _, replenishment = _copy_database(order_source, tmp_path)
    approved = _approved_order(database, replenishment["run_id"])
    monkeypatch.setenv("HACKALEM_DATA_DIR", str(database.parent))
    assert main(["order-report", "--version", str(approved["version_id"])]) == 0
    output = json.loads(capsys.readouterr().out)
    assert output["status"] == "approved"
    target = tmp_path / "exports" / "заказ.csv"
    assert main(["order-export", "--version", str(approved["version_id"]),
                 "--output", str(target)]) == 0
    exported = json.loads(capsys.readouterr().out)
    assert exported["verified"] is True and exported["path"] == str(target.resolve())
    assert target.read_bytes().startswith(b"\xef\xbb\xbf")


def test_projects_and_exports_remain_separated_by_supplier(order_source, tmp_path):
    database, _, system_run = _copy_database(order_source, tmp_path)
    system_order = create_order_project(database, system_run["run_id"], "pytest")
    iek_snapshot = next(row for row in list_snapshots(database) if row["supplier"] == "IEK")
    quality_id = list_quality_runs(database, iek_snapshot["id"])[0]["id"]
    cleaning = run_cleaning(database, iek_snapshot["id"], "2026-01-01")
    forecast = run_forecast(
        database, quality_id, cleaning["run_id"], "SYN-B-002",
        _forecast_configuration(), allow_scenario=True,
    )
    iek_run = run_replenishment(database, {
        "quality_run_id": quality_id, "cleaning_run_id": cleaning["run_id"],
        "as_of": "2026-01-01", "supplier": "IEK",
        "warehouse": "all_selected_warehouses", "items": [{
            "sku": "SYN-B-002", "category_code": "SYN-CAT-01",
            "forecast_run_id": forecast["run_id"], "project_commitments": [],
        }],
    })
    iek_order = create_order_project(database, iek_run["run_id"], "pytest")
    assert system_order["supplier"] == "Systeme Electric"
    assert iek_order["supplier"] == "IEK"
    assert {row["Поставщик"] for row in build_order_export(
        database, system_order["version_id"], "csv")["rows"]} == {"Systeme Electric"}
    assert {row["Поставщик"] for row in build_order_export(
        database, iek_order["version_id"], "csv")["rows"]} == {"IEK"}


def test_stage8_schema_migrates_to_orders_without_changing_calculations(order_source, tmp_path):
    database, _, replenishment = _copy_database(order_source, tmp_path)
    with sqlite3.connect(database) as connection:
        before = connection.execute(
            "SELECT input_json,summary_json FROM replenishment_runs WHERE id=?",
            (replenishment["run_id"],),
        ).fetchone()
        for table in ("order_events", "order_items", "order_versions", "order_projects"):
            connection.execute(f"DROP TABLE {table}")
        connection.execute("PRAGMA user_version=8")
    assert initialize_database(database).schema_version == SCHEMA_VERSION == 9
    assert initialize_database(database).schema_version == 9
    with sqlite3.connect(database) as connection:
        assert connection.execute(
            "SELECT input_json,summary_json FROM replenishment_runs WHERE id=?",
            (replenishment["run_id"],),
        ).fetchone() == before
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []


def test_export_keeps_codes_and_formula_like_text_safe(order_source, tmp_path):
    database, _, replenishment = _copy_database(order_source, tmp_path)
    draft = create_order_project(database, replenishment["run_id"], "pytest")
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE order_items SET code_1c='00123',name='=HYPERLINK(\"unsafe\")' WHERE version_id=?",
            (draft["version_id"],),
        )
    review = submit_order_for_review(database, draft["version_id"], "pytest", "Проверка текста")
    approved = approve_order(database, review["version_id"], "pytest lead", "Проверено")
    csv_export = build_order_export(database, approved["version_id"], "csv")
    row = list(csv.DictReader(io.StringIO(csv_export["content"].decode("utf-8-sig")), delimiter=";"))[0]
    assert row["Код1С"] == "'00123"
    assert row["Наименование"].startswith("'=")
    xlsx_export = build_order_export(database, approved["version_id"], "xlsx")
    book = load_workbook(io.BytesIO(xlsx_export["content"]), data_only=False)
    assert book["Заказ"]["D2"].value == "00123" and book["Заказ"]["D2"].data_type == "s"
    assert book["Заказ"]["E2"].value.startswith("=") and book["Заказ"]["E2"].data_type == "s"
