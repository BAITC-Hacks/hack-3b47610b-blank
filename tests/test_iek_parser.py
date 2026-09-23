"""Independent IEK fixtures. Partner source workbooks are never edited."""

import json
import socket
from pathlib import Path
from types import SimpleNamespace
import urllib.request
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from hackalem.importers.iek import parse_workbook
from hackalem.importers.iek_cache import append_external_cache


XML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
MOQ_HEADERS = ["№", "Код 1с", "Артикул поставщика", "Наименование", "Мин. разр. к отгр."]
TX_HEADERS = ["Дата", "Номер", "Документ", "Код", "Номенклатура", "Ед.", "Склад", "Количество"]


def _save(path: Path, rows: list[list], title: str = "Лист") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()
    return path


def _rewrite_sheet(path: Path, transform) -> None:
    with ZipFile(path) as archive:
        contents = [(info, archive.read(info.filename)) for info in archive.infolist()]
    temporary = path.with_suffix(".rewritten.xlsx")
    with ZipFile(temporary, "w") as archive:
        for info, data in contents:
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                transform(root)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            archive.writestr(info, data)
    temporary.replace(path)


def _set_cached_error(path: Path, coordinate: str) -> None:
    def transform(root):
        cell = root.find(f".//{{{XML_NS}}}c[@r='{coordinate}']")
        assert cell is not None
        cell.set("t", "e")
        value = cell.find(f"{{{XML_NS}}}v")
        if value is None:
            value = ET.SubElement(cell, f"{{{XML_NS}}}v")
        value.text = "#N/A"

    _rewrite_sheet(path, transform)


def test_minimum_errors_are_not_defaulted_and_duplicate_sku_facts_survive(tmp_path):
    path = _save(tmp_path / "minimums.xlsx", [
        MOQ_HEADERS,
        [1, "0001_", "ART-01", "Товар", 6],
        [2, "0001_", "ART-01", "Другое название того же кода", 6],
        [3, "0002_", "ART-02", "Нет значения прайса", "=VLOOKUP(C4,[1]Прайс!$A:$P,15,0)"],
    ])
    _set_cached_error(path, "E4")
    before = path.read_bytes()

    parsed = parse_workbook(path)

    assert parsed.source_kind == "minimums"
    facts = [item for item in parsed.measures if item["metric"] == "minimum_order"]
    assert [(item["row"], item["sku"], item["number"], item["state"]) for item in facts] == [
        (2, "0001_", 6, "value"),
        (3, "0001_", 6, "value"),
        (4, "0002_", None, "error"),
    ]
    assert {item["sku"] for item in parsed.products} == {"0001_", "0002_"}
    assert any(item["code"] == "DUPLICATE_SKU" for item in parsed.issues)
    raw = next(item for item in parsed.raw_rows if item["row"] == 4)
    assert raw["cells"]["E"]["value"] == "#N/A"
    assert raw["cells"]["E"]["formula"] == "=VLOOKUP(C4,[1]Прайс!$A:$P,15,0)"
    assert path.read_bytes() == before


def _stock_file(path: Path) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Остатки"
    for row in [
        ["Номенклатура", "Ед.", "Номенклатура.Код", "Январь 2024 г.", "Февраль 2024 г.", "Итого"],
        [None, None, None, "Количество", "Количество", "Количество"],
        [None, None, None, "нач. остаток", "нач. остаток", "нач. остаток"],
        ["Товар", "шт", "0001_", 10, 4, 10],
    ]:
        sheet.append(row)
    sheet.row_dimensions[3].hidden = True
    sheet.row_dimensions[6].hidden = True
    workbook.save(path)
    workbook.close()
    return path


def test_monthly_opening_stock_and_reported_total_remain_distinct(tmp_path):
    parsed = parse_workbook(_stock_file(tmp_path / "opening-stock.xlsx"))

    assert parsed.source_kind == "monthly_stock"
    assert parsed.snapshot_date is None
    assert [(item["period"], item["quantity"], item["series"]) for item in parsed.monthly_values] == [
        ("2024-01-01", 10, "opening_stock"),
        ("2024-02-01", 4, "opening_stock"),
    ]
    totals = [item for item in parsed.measures if item["metric"] == "reported_opening_stock_total"]
    assert len(totals) == 1
    assert totals[0]["number"] == 10
    assert totals[0]["cell"] == "F4"
    assert not any(item["metric"] in {"stock", "free_stock"} for item in parsed.measures)
    hidden = {(item["sheet"], item["row"]) for item in parsed.row_metadata if item["hidden"]}
    assert ("Остатки", 3) in hidden
    assert ("Остатки", 6) in hidden


def test_xml_error_type_without_value_is_blank_not_excel_error(tmp_path):
    path = _stock_file(tmp_path / "empty-error-type.xlsx")

    def transform(root):
        cell = root.find(f".//{{{XML_NS}}}c[@r='D4']")
        assert cell is not None
        cell.set("t", "e")
        for child in list(cell):
            cell.remove(child)
        zero = root.find(f".//{{{XML_NS}}}c[@r='E4']/{{{XML_NS}}}v")
        assert zero is not None
        zero.text = "0"

    _rewrite_sheet(path, transform)

    parsed = parse_workbook(path)

    assert [(item["quantity"], item["state"]) for item in parsed.monthly_values] == [
        (None, "blank"), (0, "value"),
    ]
    assert not any(item["cell"] == "D4" and item["severity"] == "error" for item in parsed.issues)


def test_incoming_order_headers_and_service_rows_keep_original_units(tmp_path):
    path = _save(tmp_path / "Путь ИЭК 22.09.2026.xlsx", [
        ["Код 1с", "Артикул ИЭК", "Наименование",
         "РФ УТ-7583 от 31 августа 2026 г. (поступление до 10.10.2026)",
         "ПП УТ-7848 от 7 сентября 2026 г. (поступление до 01.10.2026)"],
        ["0001_", "ART-01", "Товар (4шт/компл): закупаются упаковками, садятся штуками", 5, 2],
        [0, None, "Служебная строка", 999, 999],
        ["1", None, "Служебная строка", 999, 999],
        ["0001_", "ART-01", "Товар", 3, None],
    ])

    parsed = parse_workbook(path)

    assert parsed.source_kind == "incoming"
    assert {item["sku"] for item in parsed.products} == {"0001_"}
    assert [(item["cell"], item["quantity"]) for item in parsed.incoming_orders] == [
        ("D2", 5), ("E2", 2), ("D5", 3),
    ]
    first = parsed.incoming_orders[0]
    assert first["order_number"] == "УТ-7583"
    assert first["order_date"] == "2026-08-31"
    assert first["eta_deadline"] == "2026-10-10"
    assert first["header_cell"] == "D1"
    assert first["unit"] is None
    header = next(item for item in parsed.raw_rows if item["row"] == 1)
    assert header["cells"]["D"]["value"].startswith("РФ УТ-7583 от")
    second = parsed.incoming_orders[1]
    assert second["order_date"] == "2026-09-07"
    assert second["eta_deadline"] == "2026-10-01"
    assert {item["row"] for item in parsed.raw_rows}.issuperset({3, 4})


def test_identical_document_rows_are_not_deduplicated(tmp_path):
    transaction = [
        "22.09.2026 15:55:51", "000001234", "Расходная накладная",
        "0001_", "Товар", "м", "Алматы", 15,
    ]
    path = _save(tmp_path / "renamed.xlsx", [TX_HEADERS, transaction, transaction.copy()])

    parsed = parse_workbook(path)

    assert parsed.source_kind == "transactions"
    assert [(item["row"], item["sku"], item["quantity"]) for item in parsed.transactions] == [
        (2, "0001_", 15), (3, "0001_", 15),
    ]
    assert {item["document_number"] for item in parsed.transactions} == {"000001234"}


def _cache_fixture(path: Path, price_header: str, target: str) -> Path:
    _save(path, [MOQ_HEADERS, [1, "0001_", "00-A", "Товар", 6]])
    relation_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    external = ET.Element(f"{{{XML_NS}}}externalLink")
    book = ET.SubElement(external, f"{{{XML_NS}}}externalBook", {f"{{{relation_ns}}}id": "rId1"})
    names = ET.SubElement(book, f"{{{XML_NS}}}sheetNames")
    ET.SubElement(names, f"{{{XML_NS}}}sheetName", {"val": "Прайс"})
    dataset = ET.SubElement(book, f"{{{XML_NS}}}sheetDataSet")
    sheet = ET.SubElement(dataset, f"{{{XML_NS}}}sheetData", {"sheetId": "0"})
    # Columns differ deliberately from the real A:P cache: headers define meaning.
    rows = {
        7: ["Артикул", "Наименование", "Ед.", "Кратн.", "Мин. разр. к отгр.", price_header],
        8: ["00-A", "Товар", "компл", 2, 6, 1250.5],
        9: ["00-A", "Дубликат артикула", "компл", None, "#N/A", 0],
        10: ["00-B", "Другой товар", "м", 1, 1, None],
    }
    for row_number, values in rows.items():
        row = ET.SubElement(sheet, f"{{{XML_NS}}}row", {"r": str(row_number)})
        for index, value in enumerate(values, 1):
            if value is None:
                continue
            dtype = "e" if value == "#N/A" else "n" if isinstance(value, (int, float)) else "str"
            cell = ET.SubElement(row, f"{{{XML_NS}}}cell", {
                "r": f"{get_column_letter(index)}{row_number}", "t": dtype,
            })
            ET.SubElement(cell, f"{{{XML_NS}}}v").text = str(value)
    relationships = ET.Element(f"{{{package_ns}}}Relationships")
    ET.SubElement(relationships, f"{{{package_ns}}}Relationship", {
        "Id": "rId1", "Type": f"{relation_ns}/externalLinkPath",
        "Target": target, "TargetMode": "External",
    })
    with ZipFile(path) as archive:
        contents = [(info, archive.read(info.filename)) for info in archive.infolist()]
    temporary = path.with_suffix(".cached.xlsx")
    with ZipFile(temporary, "w") as archive:
        for info, data in contents:
            if info.filename == "xl/workbook.xml":
                root = ET.fromstring(data)
                refs = ET.SubElement(root, f"{{{XML_NS}}}externalReferences")
                ET.SubElement(refs, f"{{{XML_NS}}}externalReference", {f"{{{relation_ns}}}id": "rIdExternalTest"})
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename == "xl/_rels/workbook.xml.rels":
                root = ET.fromstring(data)
                ET.SubElement(root, f"{{{package_ns}}}Relationship", {
                    "Id": "rIdExternalTest", "Type": f"{relation_ns}/externalLink",
                    "Target": "/xl/externalLinks/externalLink1.xml",
                })
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename == "[Content_Types].xml":
                root = ET.fromstring(data)
                ET.SubElement(root, f"{{{content_ns}}}Override", {
                    "PartName": "/xl/externalLinks/externalLink1.xml",
                    "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
                })
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            archive.writestr(info, data)
        archive.writestr("xl/externalLinks/externalLink1.xml", ET.tostring(external, encoding="utf-8", xml_declaration=True))
        archive.writestr("xl/externalLinks/_rels/externalLink1.xml.rels", ET.tostring(relationships, encoding="utf-8", xml_declaration=True))
    temporary.replace(path)
    return path


@pytest.mark.parametrize(("price_header", "target", "currency", "price_date"), [
    ("Базовая Цена, тенге с НДС", "https://example.invalid/IEK%20price%20from%2003.08.2026_.xlsx", "KZT", "2026-08-03"),
    ("Базовая Цена", "https://example.invalid/undated-price.xlsx", None, None),
])
def test_external_cache_is_local_archived_and_does_not_invent_metadata(
    tmp_path, monkeypatch, price_header, target, currency, price_date,
):
    path = _cache_fixture(tmp_path / "cached-price.xlsx", price_header, target)
    before = path.read_bytes()

    def network_forbidden(*args, **kwargs):
        raise AssertionError("External-cache import must not access the network")

    monkeypatch.setattr(urllib.request, "urlopen", network_forbidden)
    monkeypatch.setattr(socket, "create_connection", network_forbidden)
    monkeypatch.setattr(socket.socket, "connect", network_forbidden)
    result = SimpleNamespace(catalog_items=[], external_sources=[], raw_rows=[], sheets=[], issues=[])

    append_external_cache(path, result)

    assert len(result.external_sources) == 1
    source = result.external_sources[0]
    assert source["external_target"] == target
    assert source["status"] == "archived_external_cache"
    assert source["price_date"] == price_date
    assert len(result.catalog_items) == 3
    assert [item["article"] for item in result.catalog_items] == ["00-A", "00-A", "00-B"]
    first, duplicate, last = result.catalog_items
    assert first["purchase_unit"] == "компл"
    assert first["order_multiple"] == 2
    assert first["minimum_order"] == 6
    assert first["base_price"] == 1250.5
    assert {item["currency"] for item in result.catalog_items} == {currency}
    assert {item["price_date"] for item in result.catalog_items} == {price_date}
    assert duplicate["order_multiple"] is None
    assert duplicate["minimum_order"] is None
    assert duplicate["base_price"] == 0
    states = json.loads(duplicate["states_json"])
    assert states["order_multiple"] == "blank"
    assert states["minimum_order"] == "error"
    assert states["base_price"] == "value"
    assert last["base_price"] is None
    assert json.loads(last["states_json"])["base_price"] == "blank"
    assert any(issue["code"] == "EXTERNAL_CACHE_CELL_ERROR" for issue in result.issues)
    assert any(item["state"] == "external_cache" for item in result.sheets)
    assert path.read_bytes() == before
