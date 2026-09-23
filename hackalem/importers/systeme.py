"""Read Systeme Electric reports without editing or recalculating their sources.

This module preserves alternative reports, blanks and formula caches separately.
It performs structural extraction only; reconciliation and business interpretation
belong to later stages. Source type is detected from headers, not the filename.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from itertools import zip_longest
from math import isfinite
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

RULES_VERSION = "systeme-1"


@dataclass
class ParsedWorkbook:
    source_kind: str
    snapshot_date: str | None = None
    sheets: list[dict[str, Any]] = field(default_factory=list)
    raw_rows: list[dict[str, Any]] = field(default_factory=list)
    products: list[dict[str, Any]] = field(default_factory=list)
    transactions: list[dict[str, Any]] = field(default_factory=list)
    monthly_values: list[dict[str, Any]] = field(default_factory=list)
    measures: list[dict[str, Any]] = field(default_factory=list)
    seasonal_values: list[dict[str, Any]] = field(default_factory=list)
    issues: list[dict[str, Any]] = field(default_factory=list)


_MONTHS = (
    ("янв", "январь", "января"), ("фев", "февр", "февраль", "февраля"),
    ("мар", "март", "марта"), ("апр", "апрель", "апреля"),
    ("май", "мая"), ("июн", "июнь", "июня"),
    ("июл", "июль", "июля"), ("авг", "август", "августа"),
    ("сен", "сент", "сентябрь", "сентября"),
    ("окт", "октябрь", "октября"), ("ноя", "нояб", "ноябрь", "ноября"),
    ("дек", "декабрь", "декабря"),
)
_MONTH_NUMBER = {alias: i for i, aliases in enumerate(_MONTHS, 1) for alias in aliases}


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip().casefold().replace("ё", "е")


def _iso(value: Any) -> Any:
    return value.isoformat() if isinstance(value, (datetime, date)) else value


def _formula(cell: Any) -> str | None:
    value = getattr(cell, "value", None)
    if getattr(cell, "data_type", None) != "f":
        return None
    # ArrayFormula is not a string, but its .text is the original Excel formula.
    return value if isinstance(value, str) else getattr(value, "text", str(value))


def _month(value: Any) -> str | None:
    if isinstance(value, (datetime, date)):
        return f"{value.year:04d}-{value.month:02d}-01"
    text = _normalize(value)
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:t\d{2}:\d{2}:\d{2}(?:\.\d+)?)?", text):
        try:
            parsed = datetime.fromisoformat(text)
            return f"{parsed.year:04d}-{parsed.month:02d}-01"
        except ValueError:
            return None
    match = re.fullmatch(r"([а-я]+)\.?\s+(\d{4})(?:\s*г\.?)?", text)
    if match and match[1] in _MONTH_NUMBER:
        return f"{int(match[2]):04d}-{_MONTH_NUMBER[match[1]]:02d}-01"
    return None


def _issue(result: ParsedWorkbook, code: str, message: str, *, sheet: str,
           row: int | None = None, cell: str | None = None,
           severity: str = "warning") -> None:
    result.issues.append(dict(severity=severity, code=code, sheet=sheet,
                              row=row, cell=cell, message=message))


def _detect(rows: list[dict[str, Any]]) -> dict[str, Any] | None:
    """Identify a sheet using a header row within its first twelve physical rows."""
    for record in rows:
        if record["row"] > 12:
            break
        cells = record["cells"]
        labels = {_normalize(c["value"]): col for col, c in cells.items()
                  if c["value"] is not None}
        months = {col: period for col, c in cells.items()
                  if (period := _month(c["value"])) is not None}
        has = lambda *names: all(name in labels for name in names)
        kind = None
        if has("дата", "номер", "документ", "код", "номенклатура", "ед.",
               "склад", "количество"):
            kind = "transactions"
        elif has("артикул поставщика", "код 1с", "наименование", "остаток",
                 "зарезервировано", "свободный остаток"):
            kind = "current"
        elif has("номенклатура", "номенклатура.код", "артикул", "кратность"):
            kind = "monthly_sales" if months else "multiples"
        elif has("номенклатура", "номенклатура.код", "ед.изм") and months:
            kind = "monthly_stock"
        elif "год" in labels:
            season_months = {col: _MONTH_NUMBER[_normalize(c["value"]).rstrip(".")]
                             for col, c in cells.items()
                             if _normalize(c["value"]).rstrip(".") in _MONTH_NUMBER}
            if set(season_months.values()) == set(range(1, 13)):
                return dict(kind="seasonality", row=record["row"], labels=labels,
                            months=season_months)
        if kind:
            return dict(kind=kind, row=record["row"], labels=labels, months=months)
    return None


def _cell(record: dict[str, Any], col: str | None) -> dict[str, Any]:
    if col is None:
        return {"value": None, "formula": None, "data_type": "n", "number_format": "General"}
    return record["cells"].get(col, {"value": None, "formula": None,
                                   "data_type": "n", "number_format": "General"})


def _text(record: dict[str, Any], col: str | None) -> str | None:
    value = _cell(record, col)["value"]
    if value is None:
        return None
    return str(value).strip()


def _state(cell: dict[str, Any]) -> str:
    if cell.get("cached_error") or cell["data_type"] == "e":
        return "error"
    if cell["formula"] is not None and cell["value"] is None:
        return "error"
    if cell["value"] is None or cell["value"] == "":
        return "blank"
    return "value"


def _number(record: dict[str, Any], col: str, result: ParsedWorkbook,
            *, negative_warning: bool = False) -> tuple[int | float | None, str]:
    source = _cell(record, col)
    state = _state(source)
    if state != "value":
        return None, state
    value = source["value"]
    location = f"{col}{record['row']}"
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
        _issue(result, "INVALID_NUMBER", "В числовом поле находится нечисловое значение.",
               sheet=record["sheet"], row=record["row"], cell=location, severity="error")
        return None, "error"
    if negative_warning and value < 0:
        _issue(result, "NEGATIVE_QUANTITY", "Отрицательное количество сохранено без исправления; требуется определить вид операции.",
               sheet=record["sheet"], row=record["row"], cell=location)
    return value, "value"


def _measure(result: ParsedWorkbook, record: dict[str, Any], sku: str | None,
             col: str, metric: str, *, textual: bool = False,
             negative_warning: bool = False) -> None:
    if textual:
        source = _cell(record, col)
        state = _state(source)
        number, text = None, _text(record, col) if state == "value" else None
    else:
        number, state = _number(record, col, result, negative_warning=negative_warning)
        text = None
    result.measures.append(dict(sheet=record["sheet"], row=record["row"], sku=sku,
                                metric=metric, number=number, text=text,
                                state=state, cell=f"{col}{record['row']}"))


def _parse_date(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text).isoformat()
    except ValueError:
        pass
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).isoformat()
        except ValueError:
            pass
    return None


def _snapshot(path: Path) -> str | None:
    matches = re.findall(r"(?<!\d)(\d{2})\.(\d{2})\.(\d{4})(?!\d)", path.stem)
    if len(matches) == 1:
        day, month, year = matches[0]
        try:
            return date(int(year), int(month), int(day)).isoformat()
        except ValueError:
            return None
    return None


def _read_raw(path: Path, result: ParsedWorkbook) -> dict[str, list[dict[str, Any]]]:
    groups: dict[str, list[dict[str, Any]]] = {}
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    values = None
    try:
        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        for sheet in formulas:
            cached = values[sheet.title]
            try:
                declared = sheet.calculate_dimension()
            except ValueError:
                declared = None
            # Real exports have an invalid lower bound, e.g. H77314:H77314.
            sheet.reset_dimensions()
            cached.reset_dimensions()
            rows: list[dict[str, Any]] = []
            max_row = max_column = 0
            streams = zip_longest(sheet.iter_rows(min_row=1), cached.iter_rows(min_row=1), fillvalue=())
            for row_number, (source_row, cached_row) in enumerate(streams, 1):
                max_row = row_number
                max_column = max(max_column, len(source_row), len(cached_row))
                cells: dict[str, Any] = {}
                for j, (original, value_cell) in enumerate(zip_longest(source_row, cached_row), 1):
                    formula = _formula(original)
                    value = _iso(getattr(value_cell, "value", None))
                    # Excel can cache a deliberate empty string as t="str"
                    # with an empty value element. This is distinct from an
                    # absent numeric formula cache (openpyxl returns None for both).
                    if formula is not None and value is None and getattr(value_cell, "data_type", None) == "str":
                        value = ""
                    original_value = getattr(original, "value", None)
                    if formula is None and value is None and original_value is None:
                        continue
                    # Non-formula values are copied from the source, not recomputed.
                    if formula is None and value is None:
                        value = _iso(original_value)
                    letter = get_column_letter(j)
                    item = dict(value=value, formula=formula,
                                data_type=getattr(original, "data_type", "n"),
                                number_format=getattr(original, "number_format", "General"))
                    if getattr(value_cell, "data_type", None) == "e":
                        item["cached_error"] = True
                    cells[letter] = item
                    if formula is not None and value is None:
                        _issue(result, "FORMULA_CACHE_MISSING", "У формулы нет сохранённого результата; импорт не пересчитывает Excel.",
                               sheet=sheet.title, row=row_number, cell=f"{letter}{row_number}", severity="error")
                    elif item.get("cached_error") or item["data_type"] == "e":
                        _issue(result, "EXCEL_CELL_ERROR", "В источнике сохранена ошибка Excel.",
                               sheet=sheet.title, row=row_number, cell=f"{letter}{row_number}", severity="error")
                if cells:
                    record = dict(sheet=sheet.title, row=row_number, cells=cells)
                    rows.append(record)
                    result.raw_rows.append(record)
            result.sheets.append(dict(sheet=sheet.title, state=sheet.sheet_state,
                                      declared_dimension=declared, max_row=max_row,
                                      max_column=max_column))
            groups[sheet.title] = rows
    finally:
        formulas.close()
        if values is not None:
            values.close()
    return groups


def _seasonal(result: ParsedWorkbook, rows: list[dict[str, Any]], schema: dict[str, Any]) -> None:
    year_col = schema["labels"]["год"]
    for record in rows:
        if record["row"] <= schema["row"]:
            continue
        year = _cell(record, year_col)["value"]
        if isinstance(year, bool) or not isinstance(year, (int, float)) or year != int(year) or not 1900 <= year <= 2200:
            continue
        for col, month in schema["months"].items():
            value, state = _number(record, col, result)
            result.seasonal_values.append(dict(sheet=record["sheet"], row=record["row"],
                                                period=f"{int(year):04d}-{month:02d}-01",
                                                value=value, state=state,
                                                cell=f"{col}{record['row']}"))
    _issue(result, "SEASONAL_UNIT_UNKNOWN", "Единица агрегатов сезонности не указана; исходные величины сохранены отдельно от количественных продаж.",
           sheet=rows[0]["sheet"], row=schema["row"])


def _product_rows(result: ParsedWorkbook, rows: list[dict[str, Any]], schema: dict[str, Any]):
    labels = schema["labels"]
    kind = schema["kind"]
    code_col = labels["код" if kind == "transactions" else "код 1с" if kind == "current" else "номенклатура.код"]
    name_col = labels["наименование" if kind == "current" else "номенклатура"]
    article_col = labels.get("артикул поставщика" if kind == "current" else "артикул")
    unit_col = labels.get("ед." if kind == "transactions" else "ед.изм")
    seen: dict[str, tuple[str | None, str | None, str | None]] = {}
    for record in rows:
        if record["row"] <= schema["row"]:
            continue
        # Totals are source rows, but not products. Do not drop the final stock SKU.
        if any(_normalize(_cell(record, col)["value"]) == "итого" for col in ("A", "B", name_col)):
            continue
        sku_source = _cell(record, code_col)
        sku = _text(record, code_col)
        if not sku or _state(sku_source) == "error":
            # Secondary quantity headers and completely empty rows have no name.
            if _text(record, name_col):
                _issue(result, "SKU_MISSING", "Строка товара не содержит корректного кода; исходная строка сохранена.",
                       sheet=record["sheet"], row=record["row"], cell=f"{code_col}{record['row']}", severity="error")
            continue
        if not isinstance(sku_source["value"], str):
            _issue(result, "SKU_STORED_AS_NUMBER", "Код был числом Excel и сохранён строкой; ведущие нули требуют проверки.",
                   sheet=record["sheet"], row=record["row"], cell=f"{code_col}{record['row']}")
        name, article, unit = (_text(record, col) for col in (name_col, article_col, unit_col))
        identity = name, article, unit
        if sku not in seen:
            result.products.append(dict(sheet=record["sheet"], row=record["row"], sku=sku,
                                        name=name, article=article, unit=unit,
                                        cell=f"{code_col}{record['row']}"))
            seen[sku] = identity
        elif kind != "transactions" or seen[sku] != identity:
            _issue(result, "DUPLICATE_SKU", "Повторяющийся код сохранён в фактах; справочник листа использует первое вхождение.",
                   sheet=record["sheet"], row=record["row"], cell=f"{code_col}{record['row']}")
        yield record, sku


def _transactions(result: ParsedWorkbook, rows: list[dict[str, Any]], schema: dict[str, Any]) -> None:
    labels = schema["labels"]
    sheet = rows[0]["sheet"]
    _issue(result, "CLIENT_ID_MISSING", "В документной выгрузке нет идентификатора клиента; номер документа не заменяет клиента.",
           sheet=sheet, row=schema["row"])
    start_index = len(result.transactions)
    for record, sku in _product_rows(result, rows, schema):
        qcol = labels["количество"]
        quantity, state = _number(record, qcol, result, negative_warning=True)
        if state == "blank":
            _issue(result, "QUANTITY_BLANK", "Количество отсутствует и не заменено нулём.",
                   sheet=sheet, row=record["row"], cell=f"{qcol}{record['row']}")
        occurred_at = _parse_date(_cell(record, labels["дата"])["value"])
        if occurred_at is None:
            _issue(result, "INVALID_DATE", "Дата документа отсутствует или не распознана; строка сохранена.",
                   sheet=sheet, row=record["row"], cell=f"{labels['дата']}{record['row']}", severity="error")
        document = _text(record, labels["документ"])
        document_type = re.split(r"\s+\S+\s+от\s+", document, maxsplit=1)[0] if document else None
        result.transactions.append(dict(sheet=sheet, row=record["row"], sku=sku,
                                         occurred_at=occurred_at,
                                         document_number=_text(record, labels["номер"]),
                                         document_type=document_type, unit=_text(record, labels["ед."]),
                                         warehouse=_text(record, labels["склад"]),
                                         quantity=quantity, state=state,
                                         cell=f"{qcol}{record['row']}"))
    observations = result.transactions[start_index:]
    positive_years = [int(x["occurred_at"][:4]) for x in observations
                      if x["occurred_at"] and x["quantity"] is not None and x["quantity"] > 0]
    first_year = min(positive_years) if positive_years else None
    earlier = [x for x in observations if x["occurred_at"] and x["quantity"] is not None
               and x["quantity"] < 0 and first_year is not None
               and int(x["occurred_at"][:4]) < first_year]
    if earlier:
        first = earlier[0]
        _issue(result, "OLD_NEGATIVE_TRANSACTIONS", f"Есть {len(earlier)} отрицательных строк за годы до первых положительных продаж; они не доказывают полноту ранней истории.",
               sheet=sheet, row=first["row"], cell=first["cell"])


_CURRENT_MEASURES = {
    "сс реал": "reported_cost", "витрина": "display_stock",
    "остаток тз": "sales_floor_stock", "рц ект рыскулова": "distribution_center_stock",
    "розничный склад": "retail_stock", "остаток": "stock",
    "зарезервировано": "reserved_stock", "свободный остаток": "free_stock",
    "заказ": "proposed_order", "сумма последние 12 мес": "reported_last12_total",
    "ср мес за последние 12 мес": "reported_last12_average",
    "кэф. роста": "reported_growth_rate", "кэф. сез-ти": "reported_seasonal_change",
    "запас": "reported_stock_cover", "вес": "weight",
}


def _quantities(result: ParsedWorkbook, rows: list[dict[str, Any]], schema: dict[str, Any]) -> None:
    labels, kind = schema["labels"], schema["kind"]
    sheet, header_row = rows[0]["sheet"], schema["row"]
    if kind in ("multiples", "monthly_sales"):
        _issue(result, "MULTIPLE_NOT_MOQ", "Поле «Кратность» сохранено как кратность; отдельный минимум заказа отсутствует.",
               sheet=sheet, row=header_row, cell=f"{labels['кратность']}{header_row}")
    category_cols = {label: col for label, col in labels.items() if re.fullmatch(r"категория(?:\s+\d{4})?", label)}
    incoming_cols = {label: col for label, col in labels.items() if label.startswith("сэ в пути")}
    if kind == "current":
        for label, col in category_cols.items():
            _issue(result, "CATEGORY_MEANING_UNKNOWN", "Категории источника не расшифрованы и не интерпретируются как ABC.",
                   sheet=sheet, row=header_row, cell=f"{col}{header_row}")
        if "сс реал" in labels:
            _issue(result, "REPORTED_COST_UNCONFIRMED", "«СС реал» сохранена без трактовки как подтверждённой закупочной цены или валюты.",
                   sheet=sheet, row=header_row, cell=f"{labels['сс реал']}{header_row}")
        for label, col in incoming_cols.items():
            if re.search(r"\d{1,2}\.\d{1,2}(?!\.\d{4})\b", label) and not re.search(r"\d{1,2}\.\d{1,2}\.\d{4}\b", label):
                _issue(result, "ETA_YEAR_MISSING", "Дата ожидаемого поступления в заголовке не содержит года; полная ETA не выведена из даты файла.",
                       sheet=sheet, row=header_row, cell=f"{col}{header_row}")
        if not result.snapshot_date:
            _issue(result, "SNAPSHOT_DATE_MISSING", "В имени нет однозначной полной даты текущего снимка; дата снимка не назначена.",
                   sheet=sheet, row=header_row)
    rolling_warning_emitted = False
    for record, sku in _product_rows(result, rows, schema):
        for col, period in schema["months"].items():
            quantity, state = _number(record, col, result, negative_warning=True)
            result.monthly_values.append(dict(sheet=sheet, row=record["row"], sku=sku,
                                               period=period, series="stock" if kind == "monthly_stock" else "sales",
                                               quantity=quantity, state=state,
                                               cell=f"{col}{record['row']}"))
        if kind in ("multiples", "monthly_sales"):
            col = labels["кратность"]
            _measure(result, record, sku, col, "order_multiple")
            number = result.measures[-1]["number"]
            if number is not None and (number <= 0 or number != int(number)):
                _issue(result, "INVALID_ORDER_MULTIPLE", "Кратность должна быть положительным целым числом; исходное значение сохранено без замены.",
                       sheet=sheet, row=record["row"], cell=f"{col}{record['row']}")
        if kind == "monthly_sales" and "итого" in labels:
            _measure(result, record, sku, labels["итого"], "reported_total_sales")
        if kind != "current":
            continue
        for label, metric in _CURRENT_MEASURES.items():
            if label in labels:
                _measure(result, record, sku, labels[label], metric)
        for label, col in labels.items():
            if re.fullmatch(r"продажи \d{4}", label):
                _measure(result, record, sku, col, "reported_annual_sales")
            elif re.fullmatch(r"ср мес \d{4}", label):
                _measure(result, record, sku, col, "reported_annual_average")
        for col in category_cols.values():
            _measure(result, record, sku, col, "category", textual=True)
        for label, col in incoming_cols.items():
            _measure(result, record, sku, col, "incoming_quantity", negative_warning=True)
            result.measures.append(dict(sheet=sheet, row=record["row"], sku=sku,
                                        metric="incoming_eta_label", number=None,
                                        text=_cell(next(x for x in rows if x["row"] == header_row), col)["value"],
                                        state="value", cell=f"{col}{header_row}"))
        total_col, avg_col = labels.get("сумма последние 12 мес"), labels.get("ср мес за последние 12 мес")
        total_formula = _cell(record, total_col)["formula"]
        avg_formula = _cell(record, avg_col)["formula"]
        if not rolling_warning_emitted and total_formula and avg_formula:
            match = re.fullmatch(r"=SUM\(\$?([A-Z]+)\$?\d+:\$?([A-Z]+)\$?\d+\)", total_formula.upper())
            if match and column_index_from_string(match[2]) - column_index_from_string(match[1]) + 1 == 13 and re.search(r"/12(?:\.0)?$", avg_formula):
                _issue(result, "REPORTED_12_MONTHS_USES_13", "Исходная сумма «последние 12 мес» охватывает 13 месячных колонок, среднее делит её на 12. Формулы сохранены без исправления.",
                       sheet=sheet, row=record["row"], cell=f"{total_col}{record['row']}")
                rolling_warning_emitted = True


def parse_workbook(path: Path) -> ParsedWorkbook:
    """Parse one workbook; never write sources or silently calculate missing values.

    A ValueError rejects unknown or ambiguous workbook layouts. Bad individual
    data cells remain in raw_rows and produce issues without discarding other
    products. Supplier identity is an import-context choice: transaction headers
    alone cannot distinguish a Systeme export from an IEK export.
    """
    path = Path(path)
    result = ParsedWorkbook(source_kind="")
    groups = _read_raw(path, result)
    schemas = {name: _detect(rows) for name, rows in groups.items()}
    primary = {schema["kind"] for schema in schemas.values()
               if schema is not None and schema["kind"] != "seasonality"}
    if len(primary) > 1:
        raise ValueError("В книге неоднозначный набор основных таблиц: " + ", ".join(sorted(primary)))
    if primary:
        result.source_kind = next(iter(primary))
    elif any(schema is not None for schema in schemas.values()):
        result.source_kind = "seasonality"
    else:
        raise ValueError("Структура книги Systeme Electric не распознана по заголовкам.")
    if result.source_kind == "current":
        result.snapshot_date = _snapshot(path)
    for name, rows in groups.items():
        schema = schemas[name]
        if schema is None:
            if rows:
                _issue(result, "UNRECOGNIZED_SHEET", "Лист сохранён в исходных строках, но его структура не преобразована в факты.", sheet=name)
            continue
        if schema["kind"] == "seasonality":
            _seasonal(result, rows, schema)
        elif schema["kind"] == "transactions":
            _transactions(result, rows, schema)
        else:
            _quantities(result, rows, schema)
    return result
