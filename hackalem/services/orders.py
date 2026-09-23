"""Versioned manager review, local approval and verified supplier exports."""

import csv
import io
import json
from contextlib import closing
from math import isfinite
from pathlib import Path

from openpyxl import Workbook, load_workbook

from hackalem.services.replenishment import replenishment_report
from hackalem.services.systeme import _connect, _json, _now
from hackalem.storage import initialize_database


STATUS_LABELS = {"draft": "Черновик", "review": "На проверке", "approved": "Утверждён"}
EXPORT_COLUMNS = (
    "Маркировка", "Поставщик", "Артикул", "Код1С", "Наименование",
    "Количество", "Единица закупки", "MOQ", "Кратность", "Ожидаемая дата",
    "Причина", "Проект", "Версия", "Статус",
)


def _required_text(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label} не может быть пустым.")
    return value.strip()


def _identity(connection, snapshot_id, sku):
    rows = connection.execute(
        """SELECT DISTINCT p.article,p.name FROM snapshot_files sf
        JOIN products p ON p.file_id=sf.file_id
        WHERE sf.snapshot_id=? AND p.sku=?""", (snapshot_id, sku),
    ).fetchall()
    articles = sorted({row["article"].strip() for row in rows
                       if row["article"] and row["article"].strip()})
    names = sorted({row["name"].strip() for row in rows
                    if row["name"] and row["name"].strip()})
    return {
        "article": articles[0] if len(articles) == 1 else None,
        "name": names[0] if len(names) == 1 else None,
        "article_candidates": articles,
        "name_candidates": names,
    }


def create_order_project(database_path, replenishment_run_id, created_by):
    """Create one supplier project per immutable recommendation run."""
    created_by = _required_text(created_by, "Автор проекта")
    initialize_database(Path(database_path))
    source = replenishment_report(database_path, replenishment_run_id)
    supplier = source["input"]["payload"]["supplier"]
    snapshot_id = source["input"]["snapshot_id"]
    dataset = source["input"]["dataset"]
    with closing(_connect(database_path)) as connection, connection:
        previous = connection.execute(
            "SELECT id FROM order_projects WHERE source_replenishment_run_id=?",
            (replenishment_run_id,),
        ).fetchone()
        if previous:
            version = connection.execute(
                "SELECT id FROM order_versions WHERE project_id=? ORDER BY version_number DESC LIMIT 1",
                (previous["id"],),
            ).fetchone()
            return order_report(database_path, version["id"])
        created_at = _now()
        project_id = connection.execute(
            """INSERT INTO order_projects
            (snapshot_id,supplier,source_replenishment_run_id,created_at_utc,dataset_json)
            VALUES (?,?,?,?,?)""",
            (snapshot_id, supplier, replenishment_run_id, created_at, _json(dataset)),
        ).lastrowid
        version_id = connection.execute(
            """INSERT INTO order_versions
            (project_id,version_number,parent_version_id,status,created_at_utc,created_by)
            VALUES (?,1,NULL,'draft',?,?)""", (project_id, created_at, created_by),
        ).lastrowid
        items = []
        for item in source["items"]:
            identity = _identity(connection, snapshot_id, item["sku"])
            explanation = item.get("explanation", {})
            suggested = item.get("order_quantity")
            selected = float(suggested) if suggested is not None else 0.0
            source_payload = {**item, "identity": identity}
            items.append((
                version_id, item["sku"], item["sku"], identity["article"], item["sku"],
                item.get("name") or identity["name"], suggested, selected,
                explanation.get("purchase_unit"), None, None, None, _json(source_payload),
            ))
        connection.executemany(
            """INSERT INTO order_items
            (version_id,line_id,sku,article,code_1c,name,suggested_quantity,
             selected_quantity,purchase_unit,correction_reason,changed_by,
             changed_at_utc,source_payload_json) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)""", items,
        )
        connection.execute(
            """INSERT INTO order_events
            (version_id,line_id,action,actor,reason,created_at_utc,payload_json)
            VALUES (?,NULL,'created',?,?,?,?)""",
            (version_id, created_by, "Проект создан из сохранённого расчёта.",
             created_at, _json({"replenishment_run_id": replenishment_run_id})),
        )
    return order_report(database_path, version_id)


def _version_row(connection, version_id):
    row = connection.execute(
        """SELECT v.*,p.snapshot_id,p.supplier,p.source_replenishment_run_id,
                  p.created_at_utc AS project_created_at_utc,p.dataset_json
        FROM order_versions v JOIN order_projects p ON p.id=v.project_id
        WHERE v.id=?""", (version_id,),
    ).fetchone()
    if row is None:
        raise ValueError("Версия проекта заказа не найдена.")
    return dict(row)


def order_report(database_path, version_id):
    initialize_database(Path(database_path))
    with closing(_connect(database_path)) as connection:
        result = _version_row(connection, version_id)
        result["version_id"] = result.pop("id")
        result["dataset"] = json.loads(result.pop("dataset_json"))
        result["status_label"] = STATUS_LABELS[result["status"]]
        snapshot = result.pop("approval_snapshot_json")
        result["approval_snapshot"] = json.loads(snapshot) if snapshot else None
        rows = connection.execute(
            "SELECT * FROM order_items WHERE version_id=? ORDER BY sku", (version_id,),
        )
        result["items"] = []
        for row in rows:
            item = dict(row)
            item["source"] = json.loads(item.pop("source_payload_json"))
            result["items"].append(item)
        events = connection.execute(
            "SELECT * FROM order_events WHERE version_id=? ORDER BY id", (version_id,),
        )
        result["events"] = []
        for row in events:
            event = dict(row)
            event["payload"] = json.loads(event.pop("payload_json"))
            result["events"].append(event)
        return result


def list_order_versions(database_path, snapshot_id=None):
    initialize_database(Path(database_path))
    with closing(_connect(database_path)) as connection:
        query = """SELECT v.id AS version_id,v.project_id,v.version_number,v.parent_version_id,
                   v.status,v.created_at_utc,v.approved_at_utc,v.responsible,
                   p.snapshot_id,p.supplier,p.source_replenishment_run_id,p.dataset_json
            FROM order_versions v JOIN order_projects p ON p.id=v.project_id"""
        args = []
        if snapshot_id is not None:
            query += " WHERE p.snapshot_id=?"
            args.append(snapshot_id)
        result = []
        for row in connection.execute(query + " ORDER BY v.id DESC", args):
            value = dict(row)
            value["dataset"] = json.loads(value.pop("dataset_json"))
            value["status_label"] = STATUS_LABELS[value["status"]]
            result.append(value)
        return result


def create_order_revision(database_path, approved_version_id, actor, reason):
    """Copy an approved version into a new draft without changing its snapshot."""
    actor = _required_text(actor, "Автор новой версии")
    reason = _required_text(reason, "Причина новой версии")
    initialize_database(Path(database_path))
    with closing(_connect(database_path)) as connection, connection:
        source = _version_row(connection, approved_version_id)
        if source["status"] != "approved":
            raise ValueError("Новую версию этим действием можно создать только из утверждённой.")
        number = connection.execute(
            "SELECT COALESCE(MAX(version_number),0)+1 FROM order_versions WHERE project_id=?",
            (source["project_id"],),
        ).fetchone()[0]
        created_at = _now()
        new_id = connection.execute(
            """INSERT INTO order_versions
            (project_id,version_number,parent_version_id,status,created_at_utc,created_by)
            VALUES (?,?,?,'draft',?,?)""",
            (source["project_id"], number, approved_version_id, created_at, actor),
        ).lastrowid
        connection.execute(
            """INSERT INTO order_items
            (version_id,line_id,sku,article,code_1c,name,suggested_quantity,
             selected_quantity,purchase_unit,correction_reason,changed_by,
             changed_at_utc,source_payload_json)
            SELECT ?,line_id,sku,article,code_1c,name,suggested_quantity,
                   selected_quantity,purchase_unit,correction_reason,changed_by,
                   changed_at_utc,source_payload_json
            FROM order_items WHERE version_id=?""", (new_id, approved_version_id),
        )
        connection.execute(
            """INSERT INTO order_events
            (version_id,line_id,action,actor,reason,created_at_utc,payload_json)
            VALUES (?,NULL,'revision_created',?,?,?,?)""",
            (new_id, actor, reason, created_at,
             _json({"parent_version_id": approved_version_id})),
        )
    return order_report(database_path, new_id)


def update_order_item(database_path, version_id, sku, selected_quantity, actor, reason):
    """Persist a correction by stable SKU; editing approved creates a new draft."""
    actor = _required_text(actor, "Автор корректировки")
    if (isinstance(selected_quantity, bool) or not isinstance(selected_quantity, (int, float))
            or not isfinite(selected_quantity) or selected_quantity < 0):
        raise ValueError("Выбранное количество должно быть конечным неотрицательным числом.")
    report = order_report(database_path, version_id)
    current = next((item for item in report["items"] if item["sku"] == sku), None)
    if current is None:
        raise ValueError("Строка заказа с таким SKU не найдена.")
    if float(current["selected_quantity"]) == float(selected_quantity):
        return report
    reason = _required_text(reason, "Причина корректировки")
    if report["status"] == "approved":
        report = create_order_revision(database_path, version_id, actor, reason)
        version_id = report["version_id"]
    changed_at = _now()
    with closing(_connect(database_path)) as connection, connection:
        version = _version_row(connection, version_id)
        if version["status"] == "approved":
            raise ValueError("Утверждённая версия неизменяема.")
        row = connection.execute(
            "SELECT selected_quantity FROM order_items WHERE version_id=? AND sku=?",
            (version_id, sku),
        ).fetchone()
        old_quantity = row["selected_quantity"]
        connection.execute(
            """UPDATE order_items SET selected_quantity=?,correction_reason=?,
               changed_by=?,changed_at_utc=? WHERE version_id=? AND sku=?""",
            (float(selected_quantity), reason, actor, changed_at, version_id, sku),
        )
        if version["status"] == "review":
            connection.execute(
                """UPDATE order_versions SET status='draft',submitted_at_utc=NULL
                   WHERE id=?""", (version_id,),
            )
        connection.execute(
            """INSERT INTO order_events
            (version_id,line_id,action,actor,reason,created_at_utc,payload_json)
            VALUES (?,?,'quantity_changed',?,?,?,?)""",
            (version_id, sku, actor, reason, changed_at,
             _json({"old_quantity": old_quantity, "new_quantity": float(selected_quantity)})),
        )
    return order_report(database_path, version_id)


def submit_order_for_review(database_path, version_id, actor, reason):
    actor = _required_text(actor, "Автор передачи")
    reason = _required_text(reason, "Причина передачи на проверку")
    submitted_at = _now()
    with closing(_connect(database_path)) as connection, connection:
        version = _version_row(connection, version_id)
        if version["status"] != "draft":
            raise ValueError("На проверку можно передать только черновик.")
        connection.execute(
            "UPDATE order_versions SET status='review',submitted_at_utc=? WHERE id=?",
            (submitted_at, version_id),
        )
        connection.execute(
            """INSERT INTO order_events
            (version_id,line_id,action,actor,reason,created_at_utc,payload_json)
            VALUES (?,NULL,'submitted',?,?,?,?)""",
            (version_id, actor, reason, submitted_at, _json({})),
        )
    return order_report(database_path, version_id)


def _approval_errors(report):
    positive = [item for item in report["items"] if item["selected_quantity"] > 0]
    errors = []
    if not positive:
        errors.append("В заказе нет строк с положительным количеством.")
    is_real = report["dataset"].get("kind") == "real"
    for item in positive:
        missing = [label for label, value in (
            ("артикул", item["article"]), ("код 1С", item["code_1c"]),
            ("наименование", item["name"]), ("единица закупки", item["purchase_unit"]),
        ) if value is None or (isinstance(value, str) and not value.strip())]
        if missing:
            errors.append(f"{item['sku']}: отсутствует {', '.join(missing)}.")
        source = item["source"]
        if is_real and (source.get("status") != "calculated" or source.get("scenario")):
            errors.append(f"{item['sku']}: реальная строка содержит неполные или сценарные входы.")
        explanation = source.get("explanation", {})
        if is_real and (not explanation.get("accounting_unit") or
                        not explanation.get("purchase_unit") or
                        not explanation.get("unit_factor")):
            errors.append(f"{item['sku']}: единицы или коэффициент перевода не согласованы.")
    return errors


def approve_order(database_path, version_id, responsible, note):
    """Freeze a local approval snapshot. The responsible field is not authentication."""
    responsible = _required_text(responsible, "Ответственный")
    note = _required_text(note, "Основание утверждения")
    report = order_report(database_path, version_id)
    if report["status"] != "review":
        raise ValueError("Утвердить можно только версию со статусом «На проверке».")
    errors = _approval_errors(report)
    if errors:
        raise ValueError("Утверждение заблокировано: " + " ".join(errors))
    approved_at = _now()
    source = replenishment_report(database_path, report["source_replenishment_run_id"])
    snapshot = {
        "schema": "hackalem-order-approval-1",
        "project_id": report["project_id"], "version_id": version_id,
        "version_number": report["version_number"], "status": "approved",
        "approved_at_utc": approved_at, "responsible": responsible,
        "responsible_is_authenticated_identity": False,
        "approval_note": note, "dataset": report["dataset"],
        "supplier": report["supplier"], "snapshot_id": report["snapshot_id"],
        "source_replenishment": source,
        "items": [{key: item[key] for key in (
            "line_id", "sku", "article", "code_1c", "name", "suggested_quantity",
            "selected_quantity", "purchase_unit", "correction_reason", "changed_by",
            "changed_at_utc", "source",
        )} for item in report["items"]],
    }
    with closing(_connect(database_path)) as connection, connection:
        cursor = connection.execute(
            """UPDATE order_versions SET status='approved',approved_at_utc=?,responsible=?,
               approval_note=?,approval_snapshot_json=? WHERE id=? AND status='review'
               AND approval_snapshot_json IS NULL""",
            (approved_at, responsible, note, _json(snapshot), version_id),
        )
        if cursor.rowcount != 1:
            raise ValueError("Версия уже изменена; обновите проект перед утверждением.")
        connection.execute(
            """INSERT INTO order_events
            (version_id,line_id,action,actor,reason,created_at_utc,payload_json)
            VALUES (?,NULL,'approved',?,?,?,?)""",
            (version_id, responsible, note, approved_at,
             _json({"responsible_is_authenticated_identity": False})),
        )
    return order_report(database_path, version_id)


def _export_source(report):
    if report["status"] == "approved":
        snapshot = report["approval_snapshot"]
        return snapshot["items"], snapshot
    return report["items"], {
        "project_id": report["project_id"], "version_number": report["version_number"],
        "status": report["status"], "supplier": report["supplier"],
        "dataset": report["dataset"], "approved_at_utc": None,
    }


def _classification(report, items, metadata):
    if metadata["dataset"].get("kind") == "synthetic":
        return "SYNTHETIC_SCENARIO"
    if any(item["source"].get("scenario") for item in items):
        return "SCENARIO"
    if report["status"] != "approved":
        return "DRAFT"
    return "APPROVED_REAL"


def _csv_text(value):
    if value is None:
        return ""
    text = str(value)
    if text.startswith(("=", "+", "-", "@")) or (len(text) > 1 and text[0] == "0" and text.isdigit()):
        return "'" + text
    return text


def _export_rows(report):
    items, metadata = _export_source(report)
    classification = _classification(report, items, metadata)
    supplier = metadata["supplier"]
    rows = []
    for item in items:
        if item["selected_quantity"] <= 0:
            continue
        explanation = item["source"].get("explanation", {})
        reason = item.get("correction_reason") or explanation.get("text") or "; ".join(item["source"].get("reasons", []))
        rows.append({
            "Маркировка": classification, "Поставщик": supplier,
            "Артикул": item.get("article") or "", "Код1С": item["code_1c"],
            "Наименование": item.get("name") or "", "Количество": item["selected_quantity"],
            "Единица закупки": item.get("purchase_unit") or "",
            "MOQ": explanation.get("minimum_order"), "Кратность": explanation.get("order_multiple"),
            "Ожидаемая дата": explanation.get("new_order_eta") or "",
            "Причина": reason or "", "Проект": report["project_id"],
            "Версия": report["version_number"], "Статус": report["status_label"],
        })
    return rows, {**metadata, "classification": classification, "row_count": len(rows)}


def _verify_rows(actual, expected):
    if len(actual) != len(expected):
        raise RuntimeError("Проверка экспорта: число строк изменилось.")
    for index, (left, right) in enumerate(zip(actual, expected), start=1):
        for column in EXPORT_COLUMNS:
            if column in ("Количество", "MOQ", "Кратность", "Проект", "Версия"):
                left_value = None if left[column] in (None, "") else float(left[column])
                right_value = None if right[column] in (None, "") else float(right[column])
                if left_value != right_value:
                    raise RuntimeError(f"Проверка экспорта: строка {index}, поле {column} изменилось.")
            elif str(left[column] or "") != str(right[column] or ""):
                raise RuntimeError(f"Проверка экспорта: строка {index}, поле {column} изменилось.")


def build_order_export(database_path, version_id, file_format):
    """Build and immediately re-read CSV/XLSX bytes before returning them."""
    if file_format not in {"csv", "xlsx"}:
        raise ValueError("Формат экспорта должен быть csv или xlsx.")
    report = order_report(database_path, version_id)
    rows, metadata = _export_rows(report)
    if not rows:
        raise ValueError("В выбранной версии нет положительных строк для экспорта.")
    safe_rows = [{column: (_csv_text(row[column]) if column not in
                           ("Количество", "MOQ", "Кратность", "Проект", "Версия") else row[column])
                  for column in EXPORT_COLUMNS} for row in rows]
    if file_format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.DictWriter(stream, fieldnames=EXPORT_COLUMNS, delimiter=";", lineterminator="\n")
        writer.writeheader()
        writer.writerows(safe_rows)
        content = stream.getvalue().encode("utf-8-sig")
        reader = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
        _verify_rows(reader, safe_rows)
        media_type = "text/csv"
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Заказ"
        sheet.append(EXPORT_COLUMNS)
        text_columns = {index + 1 for index, column in enumerate(EXPORT_COLUMNS)
                        if column not in ("Количество", "MOQ", "Кратность", "Проект", "Версия")}
        for row in rows:
            sheet.append([row[column] for column in EXPORT_COLUMNS])
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column in text_columns:
                    cell.number_format = "@"
                    if cell.value is not None:
                        cell.value = str(cell.value)
                        cell.data_type = "s"
        meta = workbook.create_sheet("Метаданные")
        meta.append(("Параметр", "Значение"))
        for key, value in sorted(metadata.items()):
            meta.append((key, _json(value) if isinstance(value, (dict, list)) else value))
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        checked = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        values = list(checked["Заказ"].iter_rows(values_only=True))
        actual = [dict(zip(EXPORT_COLUMNS, row)) for row in values[1:]]
        _verify_rows(actual, rows)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    filename = f"order-{report['project_id']}-v{report['version_number']}-{metadata['classification'].lower()}.{file_format}"
    return {"content": content, "filename": filename, "media_type": media_type,
            "rows": rows, "metadata": metadata, "verified": True}


def export_order_file(database_path, version_id, output_path):
    """Write an export, re-read the saved file and verify it byte-for-byte logically."""
    output_path = Path(output_path)
    file_format = output_path.suffix.lower().lstrip(".")
    result = build_order_export(database_path, version_id, file_format)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(result["content"])
    saved = output_path.read_bytes()
    if saved != result["content"]:
        raise RuntimeError("Сохранённый экспорт отличается от проверенного содержимого.")
    return {key: value for key, value in result.items() if key != "content"} | {
        "path": str(output_path.resolve()), "size": len(saved),
    }
